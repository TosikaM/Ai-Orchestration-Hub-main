#!/usr/bin/env python3
"""
Format-specific extractors for the knowledge pipeline.

Separated from ingest-knowledge.py because the *pipeline* (inbox -> extracted ->
distilled -> index) and the *readers* change for different reasons: a new client format
adds a reader here and touches nothing else.

THE RULE EVERY READER FOLLOWS
-----------------------------
Preserve the citation anchor. A PDF is cited by page, a deck by slide, a workbook by
sheet. Each reader emits an HTML comment marker — `<!-- page: 14 -->`, `<!-- slide: 7 -->`,
`<!-- sheet: Features -->` — so a distilled claim can point back at the exact place in the
original. An extraction that drops those markers is text, not evidence, and the whole
value of the corpus is that a requirement can be traced rather than recalled.

AND THE ONE THAT MATTERS MOST
-----------------------------
A reader that cannot genuinely read something returns None with a reason. It never
returns empty text that looks like a successful extraction. A scanned PDF ingested as an
empty document is worse than one reported as needing OCR, because the gap stops being
visible the moment it looks processed.
"""
from __future__ import annotations

from pathlib import Path

TEXT_SUFFIXES = {".md", ".txt", ".csv", ".json", ".yaml", ".yml", ".rst"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}
HTML_SUFFIXES = {".html", ".htm"}
OFFICE_SUFFIXES = {".pdf", ".docx", ".xlsx", ".pptx"}
LEGACY_SUFFIXES = {".doc", ".xls", ".ppt"}

LEGACY_ADVICE = {".doc": ".docx", ".xls": ".xlsx", ".ppt": ".pptx"}


def extract_pdf(path: Path):
    try:
        from pypdf import PdfReader
    except ImportError:
        return None, "pypdf is not installed. `pip install pypdf`."
    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        return None, f"could not open: {exc}"

    chunks, empty = [], 0
    for number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if not text.strip():
            empty += 1
        chunks.append(f"\n<!-- page: {number} -->\n{text}")

    if empty == len(reader.pages):
        return None, (
            f"all {len(reader.pages)} pages have no text layer — this is a scanned "
            f"document and needs OCR. Not ingesting an empty file that would look processed."
        )
    note = f"{empty} of {len(reader.pages)} pages had no text layer" if empty else None
    return "".join(chunks), note


def extract_docx(path: Path):
    try:
        import docx
    except ImportError:
        return None, "python-docx is not installed. `pip install python-docx`."
    try:
        document = docx.Document(str(path))
    except Exception as exc:
        return None, f"could not open: {exc}"

    lines = []
    for para in document.paragraphs:
        style = (para.style.name or "").lower()
        text = para.text.strip()
        if not text:
            continue
        if style.startswith("heading"):
            digits = "".join(ch for ch in style if ch.isdigit()) or "2"
            lines.append(f"\n{'#' * min(int(digits) + 1, 6)} {text}\n")
        else:
            lines.append(text)

    for index, table in enumerate(document.tables, start=1):
        lines.append(f"\n<!-- table: {index} -->\n")
        for row_index, row in enumerate(table.rows):
            cells = [c.text.strip().replace("|", "\\|") for c in row.cells]
            lines.append("| " + " | ".join(cells) + " |")
            if row_index == 0:
                lines.append("|" + "|".join(["---"] * len(cells)) + "|")

    if not lines:
        return None, "no text or tables found"
    return "\n".join(lines), None


def extract_xlsx(path: Path):
    """
    Spreadsheets are where feature lists and requirement matrices live, and they are the
    format most often mangled in ingestion because a naive reader flattens them into a
    wall of values with no headings.

    The header row is preserved deliberately: in a feature list the heading *is* the
    meaning. A cell reading "P1" is noise until you know its column is Priority.
    """
    try:
        from openpyxl import load_workbook
    except ImportError:
        return None, "openpyxl is not installed. `pip install openpyxl`."
    try:
        book = load_workbook(str(path), data_only=True, read_only=True)
    except Exception as exc:
        return None, f"could not open: {exc}"

    chunks, empty_sheets = [], 0
    for sheet in book.worksheets:
        rows = [
            ["" if cell is None else str(cell).strip().replace("|", "\\|") for cell in row]
            for row in sheet.iter_rows(values_only=True)
        ]
        rows = [r for r in rows if any(cell for cell in r)]
        if not rows:
            empty_sheets += 1
            continue
        width = max(len(r) for r in rows)
        rows = [r + [""] * (width - len(r)) for r in rows]

        chunks.append(f"\n<!-- sheet: {sheet.title} -->\n")
        chunks.append(f"## Sheet: {sheet.title}\n")
        chunks.append("| " + " | ".join(rows[0]) + " |")
        chunks.append("|" + "|".join(["---"] * width) + "|")
        for row in rows[1:]:
            chunks.append("| " + " | ".join(row) + " |")
        chunks.append("")
    book.close()

    if not chunks:
        return None, f"every sheet is empty ({empty_sheets} sheet(s))"
    note = f"{empty_sheets} empty sheet(s) skipped" if empty_sheets else None
    return "\n".join(chunks), note


def extract_pptx(path: Path):
    """
    Approach and solution-design decks. Slide number is the citation anchor, because a
    deck is cited by slide the way a document is cited by page.

    Speaker notes are extracted and marked as notes. They routinely carry the caveat the
    slide omits — which is exactly the kind of thing a requirement loses on its way to a
    backlog item.
    """
    try:
        from pptx import Presentation
    except ImportError:
        return None, "python-pptx is not installed. `pip install python-pptx`."
    try:
        deck = Presentation(str(path))
    except Exception as exc:
        return None, f"could not open: {exc}"

    chunks, empty, total = [], 0, 0
    for number, slide in enumerate(deck.slides, start=1):
        total += 1
        lines = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if text:
                        lines.append(f"{'  ' * getattr(para, 'level', 0)}- {text}")
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [c.text.strip().replace("|", "\\|") for c in row.cells]
                    lines.append("| " + " | ".join(cells) + " |")

        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()

        if not lines and not notes:
            empty += 1
            continue
        chunks.append(f"\n<!-- slide: {number} -->\n")
        chunks.append(f"## Slide {number}\n")
        chunks.extend(lines)
        if notes:
            chunks.append(f"\n> **Speaker notes:** {notes}\n")

    if not chunks:
        return None, (
            f"no text on any of {total} slide(s) — likely an image-only deck. "
            f"Export the slides as images and let a model describe them, or add OCR."
        )
    note = f"{empty} slide(s) had no text" if empty else None
    return "\n".join(chunks), note


def extract_html(path: Path):
    """
    HTML mockups. Script and style are dropped; structure is kept, because for a screen
    the structure IS the requirement — which fields exist, how they are labelled, which
    are required, what the buttons say.
    """
    try:
        from bs4 import BeautifulSoup
    except ImportError:
        return None, "beautifulsoup4 is not installed. `pip install beautifulsoup4 lxml`."
    try:
        soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "html.parser")
    except Exception as exc:
        return None, f"could not parse: {exc}"

    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()

    title = soup.title.get_text(strip=True) if soup.title else "(untitled screen)"
    out = [f"# Screen: {title}", "", f"<!-- source: {path.name} -->", ""]

    # The interactive surface first. A mockup's value to a requirements reader is the
    # field inventory, and it is buried if you only take the visible text.
    controls = []
    for el in soup.find_all(["input", "select", "textarea", "button"]):
        kind = el.name if el.name != "input" else f"input[{el.get('type', 'text')}]"
        label = (el.get("name") or el.get("id") or el.get("placeholder")
                 or el.get_text(strip=True) or "(unlabelled)")
        flags = [f for f in ("required", "disabled", "readonly", "checked") if el.has_attr(f)]
        controls.append(f"| {kind} | {str(label).replace('|', '\\|')} | {', '.join(flags) or '-'} |")
    if controls:
        out += ["## Interactive controls", "",
                "| element | label / name | flags |", "|---|---|---|", *controls, ""]

    links = [f"- `{a.get('href')}` — {a.get_text(strip=True) or '(no text)'}"
             for a in soup.find_all("a", href=True)][:40]
    if links:
        out += ["## Navigation", "", *links, ""]

    headings = [f"{'#' * min(int(h.name[1]) + 1, 6)} {h.get_text(strip=True)}"
                for h in soup.find_all(["h1", "h2", "h3", "h4"])]
    if headings:
        out += ["## Headings", "", *headings, ""]

    text = "\n".join(line.strip() for line in soup.get_text("\n").splitlines() if line.strip())
    out += ["## Visible text", "", text]
    return "\n".join(out), None


def inventory_image(path: Path):
    """
    An image is NOT extracted here, and saying so plainly is the whole point.

    A screenshot of a screen contains the requirement — the fields, the labels, the
    states, the navigation — and none of it is recoverable by a script. It needs a model
    that can see. So this stage writes a STUB recording the file and its dimensions, and
    marks it as awaiting description.

    `ingest-knowledge.py worklist` then prints exactly which images are still outstanding.
    That is the honest shape: an image silently "ingested" as an empty document is worse
    than one listed as unfinished, because the gap stops being visible the moment it looks
    processed.
    """
    try:
        from PIL import Image
        with Image.open(path) as img:
            size, mode = f"{img.width}x{img.height}", img.mode
    except Exception as exc:
        size, mode = "unknown", f"unreadable ({exc})"

    body = f"""# Screen: {path.stem}

<!-- image: {path.name} · {size} · {mode} -->
<!-- status: awaiting-description -->

**STATUS: awaiting description. This is a stub, not an extraction.**

A script cannot read a screen. The fields, labels, states and navigation in this image
are the requirement, and recovering them needs a model that can see it.

To complete it, open the image and replace everything below with:

- **Purpose** — what a user accomplishes on this screen
- **Fields** — name, type, required, validation, default value
- **Actions** — every button and control, and where each one leads
- **States** — empty, loading, error, success, permission-denied
- **Navigation** — what reaches this screen, and what it reaches
- **Open questions** — anything the image does not settle

Source: `{path.name}`
"""
    return body, "image stub — needs a model to describe it"


def dispatch(path: Path):
    """
    Route a file to its reader.
    @returns (text, note) — text is None when the file could not genuinely be read.
    """
    suffix = path.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        return path.read_text(encoding="utf-8", errors="replace"), None
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix == ".docx":
        return extract_docx(path)
    if suffix == ".xlsx":
        return extract_xlsx(path)
    if suffix == ".pptx":
        return extract_pptx(path)
    if suffix in HTML_SUFFIXES:
        return extract_html(path)
    if suffix in IMAGE_SUFFIXES:
        return inventory_image(path)
    if suffix in LEGACY_SUFFIXES:
        return None, (
            f"{suffix} is a legacy binary format with no reader here. "
            f"Re-save as {LEGACY_ADVICE[suffix]} and re-run — the conversion is lossless "
            f"for text and the modern format is readable without Office installed."
        )
    return None, f"unknown type {suffix}"
